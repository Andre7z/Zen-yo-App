from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.atividades import controle_faixas


def _aluno(nome, proxima, tamanho, grossa=False):
    return SimpleNamespace(
        nome=nome,
        proxima_faixa=proxima,
        tamanho_faixa=tamanho,
        faixa_grossa=grossa,
        faixa_obj=SimpleNamespace(nome="Branca"),
    )


def _atividade():
    alunos = [
        _aluno("Ana", "Branca/CINZA - 11º Kyu", "pp"),
        _aluno("Beto", "BRANCA/CINZA", "PP", True),
        _aluno("Caio", "Azul - 8º Kyu", "N2"),
        _aluno("Duda", None, "GG"),
    ]
    return SimpleNamespace(nova_faixa="Verde", alunos=alunos), alunos


def test_quantidades_por_faixa_e_tamanho_com_totais():
    atividade, alunos = _atividade()

    relatorio = controle_faixas.gerar_relatorio(atividade, alunos)
    quantidades = relatorio["quantidades"]
    por_faixa = {f["faixa"]: f for f in quantidades["faixas"]}

    assert relatorio["linhas"][0]["faixa"] == "BRANCA PONTA CINZA"
    assert relatorio["linhas"][1]["grossa"] is True
    assert por_faixa["BRANCA PONTA CINZA"]["tamanhos"][0] == 2
    assert por_faixa["BRANCA PONTA CINZA"]["total"] == 2
    assert por_faixa["AZUL"]["tamanhos"][2] == 1
    assert por_faixa["VERDE"]["tamanhos"][6] == 1
    assert quantidades["totais"] == [2, 0, 1, 0, 0, 0, 1]
    assert quantidades["total_geral"] == 4


def test_xlsx_preserva_layout_da_tabela_original():
    atividade, alunos = _atividade()

    planilha = load_workbook(BytesIO(controle_faixas.gerar_xlsx(atividade, alunos)))
    folha = planilha.active

    assert folha["A1"].value == "CONTROLE DE FAIXAS"
    assert "A1:M1" in [str(m) for m in folha.merged_cells.ranges]
    assert [c.value for c in folha[2]] == [
        "NOME", "FAIXA", "TAMANHO", "GROSSA", "QUANTIDADES",
        "PP", "P", "N2", "M", "G", "N4", "GG", "Total",
    ]
    assert [folha.cell(row=r, column=5).value for r in range(3, 15)] == [*controle_faixas.FAIXAS, "TOTAL:"]
    assert folha["A3"].value == "Ana"
    assert folha["D4"].value == "SIM"
    assert folha["F3"].value.startswith("=COUNTIFS(")
    assert folha["M14"].value == "=SUM(M3:M13)"
    assert folha["E3"].fill.fgColor.rgb.endswith("FFFF00")
    assert folha["B3"].fill.fgColor.rgb.endswith("DDEBF7")
    assert folha["E3"].border.left.style == "thin"
