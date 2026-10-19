"""Relatorio "CONTROLE DE FREQUENCIA" da atividade de frequencia.

Monta as linhas (nome, presenca, data) e a tabela de quantidades de
presentes/ausentes, e gera a planilha .xlsx no mesmo estilo dos demais
relatorios de atividade.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ._estilo_xlsx import AMARELO, AZUL, BORDA, CENTRO, ESQUERDA, NEGRITO, NORMAL

TITULO = "CONTROLE DE FREQUENCIA"
LINHAS_MINIMAS = 14


def montar_linhas(atividade, alunos):
    presentes = set(getattr(atividade, "aluno_ids_pagos", []) or [])
    return [
        {
            "nome": aluno.nome,
            "presenca": aluno.id in presentes,
            "data": str(atividade.data) if atividade.data else "",
        }
        for aluno in alunos
    ]


def calcular_quantidades(linhas):
    presentes = sum(1 for linha in linhas if linha["presenca"])
    ausentes = len(linhas) - presentes
    return {"presentes": presentes, "ausentes": ausentes, "total_geral": presentes + ausentes}


def gerar_relatorio(atividade, alunos):
    linhas = montar_linhas(atividade, alunos)
    return {"titulo": TITULO, "linhas": linhas, "quantidades": calcular_quantidades(linhas)}


def gerar_xlsx(atividade, alunos):
    linhas = montar_linhas(atividade, alunos)
    ultima_lista = 2 + max(len(linhas), LINHAS_MINIMAS)

    planilha = Workbook()
    folha = planilha.active
    folha.title = "Controle de frequencia"

    colunas_totais = 5
    ultima_coluna = get_column_letter(colunas_totais)
    folha.merge_cells(f"A1:{ultima_coluna}1")
    folha["A1"] = TITULO
    folha["A1"].font = NEGRITO
    folha["A1"].alignment = CENTRO
    for coluna in range(1, colunas_totais + 1):
        folha.cell(row=1, column=coluna).border = BORDA

    cabecalhos = ["NOME", "PRESENÇA", "DATA", "QUANTIDADES", "Total"]
    for indice, texto in enumerate(cabecalhos, start=1):
        celula = folha.cell(row=2, column=indice, value=texto)
        celula.font = NEGRITO
        celula.alignment = CENTRO
        celula.border = BORDA

    for numero in range(3, ultima_lista + 1):
        for coluna in range(1, 4):
            celula = folha.cell(row=numero, column=coluna)
            celula.border = BORDA
            celula.font = NORMAL
            celula.alignment = ESQUERDA if coluna == 1 else CENTRO
            if coluna in (2, 3):
                celula.fill = AZUL
    for indice, linha in enumerate(linhas):
        numero = 3 + indice
        folha.cell(row=numero, column=1, value=linha["nome"])
        folha.cell(row=numero, column=2, value="PRESENTE" if linha["presenca"] else "AUSENTE")
        folha.cell(row=numero, column=3, value=linha["data"] or None)

    intervalo_presenca = f"$B$3:$B${ultima_lista}"
    rotulos = [("PRESENTES", "PRESENTE"), ("AUSENTES", "AUSENTE")]
    for indice, (rotulo, valor) in enumerate(rotulos):
        numero = 3 + indice
        folha.cell(row=numero, column=4, value=rotulo)
        folha.cell(row=numero, column=5, value=f'=COUNTIF({intervalo_presenca},"{valor}")')

    total_linha = 3 + len(rotulos)
    folha.cell(row=total_linha, column=4, value="TOTAL:")
    folha.cell(row=total_linha, column=5, value=f"=SUM(E3:E{total_linha - 1})")

    for numero in range(3, total_linha + 1):
        for coluna in (4, 5):
            celula = folha.cell(row=numero, column=coluna)
            celula.fill = AMARELO
            celula.border = BORDA
            celula.font = NEGRITO if numero == total_linha else NORMAL
            celula.alignment = ESQUERDA if coluna == 4 else CENTRO

    larguras = {1: 42, 2: 14, 3: 14, 4: 20}
    for coluna in range(1, colunas_totais + 1):
        folha.column_dimensions[get_column_letter(coluna)].width = larguras.get(coluna, 9)

    saida = BytesIO()
    planilha.save(saida)
    return saida.getvalue()
