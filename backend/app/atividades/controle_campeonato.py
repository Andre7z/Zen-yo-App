"""Relatorio "CONTROLE DE CAMPEONATO" da atividade de campeonato.

Monta as linhas (nome, idade, peso, codigo zempo) e gera a planilha .xlsx
no mesmo estilo dos demais relatorios de atividade.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ._estilo_xlsx import AZUL, BORDA, CENTRO, ESQUERDA, NEGRITO, NORMAL

TITULO = "CONTROLE DE CAMPEONATO"
LINHAS_MINIMAS = 14


def montar_linhas(alunos):
    return [
        {
            "nome": aluno.nome,
            "idade": getattr(aluno, "idade", None),
            "peso": getattr(aluno, "peso", None),
            "codigo_zempo": getattr(aluno, "codigo_zempo", None) or "",
        }
        for aluno in alunos
    ]


def gerar_relatorio(alunos):
    linhas = montar_linhas(alunos)
    return {"titulo": TITULO, "linhas": linhas, "total_geral": len(linhas)}


def gerar_xlsx(alunos):
    linhas = montar_linhas(alunos)
    ultima_lista = 2 + max(len(linhas), LINHAS_MINIMAS)

    planilha = Workbook()
    folha = planilha.active
    folha.title = "Controle de campeonato"

    colunas_totais = 4
    folha.merge_cells(f"A1:{get_column_letter(colunas_totais)}1")
    folha["A1"] = TITULO
    folha["A1"].font = NEGRITO
    folha["A1"].alignment = CENTRO
    for coluna in range(1, colunas_totais + 1):
        folha.cell(row=1, column=coluna).border = BORDA

    for indice, texto in enumerate(["NOME", "IDADE", "PESO", "CÓDIGO ZEMPO"], start=1):
        celula = folha.cell(row=2, column=indice, value=texto)
        celula.font = NEGRITO
        celula.alignment = CENTRO
        celula.border = BORDA

    for numero in range(3, ultima_lista + 1):
        for coluna in range(1, colunas_totais + 1):
            celula = folha.cell(row=numero, column=coluna)
            celula.border = BORDA
            celula.font = NORMAL
            celula.alignment = ESQUERDA if coluna == 1 else CENTRO
            if coluna > 1:
                celula.fill = AZUL
    for indice, linha in enumerate(linhas):
        numero = 3 + indice
        folha.cell(row=numero, column=1, value=linha["nome"])
        folha.cell(row=numero, column=2, value=linha["idade"])
        folha.cell(row=numero, column=3, value=linha["peso"])
        folha.cell(row=numero, column=4, value=linha["codigo_zempo"] or None)

    larguras = {1: 42, 2: 10, 3: 12, 4: 20}
    for coluna, largura in larguras.items():
        folha.column_dimensions[get_column_letter(coluna)].width = largura

    saida = BytesIO()
    planilha.save(saida)
    return saida.getvalue()
