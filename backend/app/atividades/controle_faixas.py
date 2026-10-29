"""Relatorio "CONTROLE DE FAIXAS" da atividade de troca de faixa.

Monta as linhas (nome, faixa, tamanho, grossa), a tabela de quantidades por
faixa e tamanho e a planilha .xlsx no layout da tabela original.
"""
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITULO = "CONTROLE DE FAIXAS"
FAIXAS = [
    "BRANCA PONTA CINZA",
    "CINZA",
    "CINZA PONTA AZUL",
    "AZUL",
    "AZUL PONTA AMARELA",
    "AMARELA",
    "AMARELA PONTA LARANJA",
    "LARANJA",
    "VERDE",
    "ROXA",
    "MARROM",
]
TAMANHOS_PADRAO = ["PP", "P", "N2", "M", "G", "N4", "GG"]
LINHAS_MINIMAS = len(FAIXAS) + 3


def normalizar_faixa(texto):
    """'Branca/CINZA - 11º Kyu' -> 'BRANCA PONTA CINZA'."""
    if not texto:
        return ""
    nome = str(texto).split(" - ")[0].upper().replace("/", " PONTA ")
    return re.sub(r"\s+", " ", nome).strip()


def normalizar_tamanho(texto):
    return str(texto).strip().upper() if texto else ""


def montar_linhas(atividade, alunos):
    linhas = []
    for aluno in alunos:
        faixa = aluno.proxima_faixa or atividade.nova_faixa or aluno.faixa_obj.nome
        linhas.append(
            {
                "nome": aluno.nome,
                "faixa": normalizar_faixa(faixa),
                "tamanho": normalizar_tamanho(aluno.tamanho_faixa),
                "grossa": bool(getattr(aluno, "faixa_grossa", False)),
            }
        )
    return linhas


def calcular_quantidades(linhas, tamanhos=None):
    TAMANHOS = tamanhos or TAMANHOS_PADRAO
    matriz = {faixa: {tamanho: 0 for tamanho in TAMANHOS} for faixa in FAIXAS}
    for linha in linhas:
        if linha["faixa"] in matriz and linha["tamanho"] in TAMANHOS:
            matriz[linha["faixa"]][linha["tamanho"]] += 1
    por_faixa = [
        {
            "faixa": faixa,
            "tamanhos": [matriz[faixa][t] for t in TAMANHOS],
            "total": sum(matriz[faixa].values()),
        }
        for faixa in FAIXAS
    ]
    totais = [sum(matriz[f][t] for f in FAIXAS) for t in TAMANHOS]
    return {
        "tamanhos": TAMANHOS,
        "faixas": por_faixa,
        "totais": totais,
        "total_geral": sum(totais),
    }


def gerar_relatorio(atividade, alunos, tamanhos=None):
    linhas = montar_linhas(atividade, alunos)
    return {"titulo": TITULO, "linhas": linhas, "quantidades": calcular_quantidades(linhas, tamanhos)}


def gerar_xlsx(atividade, alunos, tamanhos=None):
    TAMANHOS = tamanhos or TAMANHOS_PADRAO
    linhas = montar_linhas(atividade, alunos)
    ultima_lista = 2 + max(len(linhas), LINHAS_MINIMAS)
    total_linha = 3 + len(FAIXAS)

    planilha = Workbook()
    folha = planilha.active
    folha.title = "Controle de faixas"

    fino = Side(style="thin", color="000000")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    azul = PatternFill("solid", fgColor="DDEBF7")
    amarelo = PatternFill("solid", fgColor="FFFF00")
    negrito = Font(name="Calibri", size=11, bold=True)
    normal = Font(name="Calibri", size=11)
    centro = Alignment(horizontal="center", vertical="center")
    esquerda = Alignment(horizontal="left", vertical="center")

    colunas_totais = 5 + len(TAMANHOS) + 1
    ultima_coluna = get_column_letter(colunas_totais)
    folha.merge_cells(f"A1:{ultima_coluna}1")
    folha["A1"] = TITULO
    folha["A1"].font = negrito
    folha["A1"].alignment = centro
    for coluna in range(1, colunas_totais + 1):
        folha.cell(row=1, column=coluna).border = borda

    cabecalhos = ["NOME", "FAIXA", "TAMANHO", "GROSSA", "QUANTIDADES", *TAMANHOS, "Total"]
    for indice, texto in enumerate(cabecalhos, start=1):
        celula = folha.cell(row=2, column=indice, value=texto)
        celula.font = negrito
        celula.alignment = centro
        celula.border = borda

    for numero in range(3, ultima_lista + 1):
        for coluna in range(1, 5):
            celula = folha.cell(row=numero, column=coluna)
            celula.border = borda
            celula.font = normal
            celula.alignment = esquerda if coluna == 1 else centro
            if coluna in (2, 3):
                celula.fill = azul
    for indice, linha in enumerate(linhas):
        numero = 3 + indice
        folha.cell(row=numero, column=1, value=linha["nome"])
        folha.cell(row=numero, column=2, value=linha["faixa"] or None)
        folha.cell(row=numero, column=3, value=linha["tamanho"] or None)
        folha.cell(row=numero, column=4, value="SIM" if linha["grossa"] else "NÃO")

    intervalo_faixa = f"$B$3:$B${ultima_lista}"
    intervalo_tamanho = f"$C$3:$C${ultima_lista}"
    for indice, faixa in enumerate(FAIXAS):
        numero = 3 + indice
        folha.cell(row=numero, column=5, value=faixa)
        for posicao, tamanho in enumerate(TAMANHOS):
            coluna = 6 + posicao
            folha.cell(
                row=numero,
                column=coluna,
                value=f'=COUNTIFS({intervalo_faixa},$E{numero},{intervalo_tamanho},{get_column_letter(coluna)}$2)',
            )
        primeira, ultima = get_column_letter(6), get_column_letter(5 + len(TAMANHOS))
        folha.cell(row=numero, column=colunas_totais, value=f"=SUM({primeira}{numero}:{ultima}{numero})")

    folha.cell(row=total_linha, column=5, value="TOTAL:")
    for coluna in range(6, colunas_totais + 1):
        letra = get_column_letter(coluna)
        folha.cell(row=total_linha, column=coluna, value=f"=SUM({letra}3:{letra}{total_linha - 1})")

    for numero in range(3, total_linha + 1):
        for coluna in range(5, colunas_totais + 1):
            celula = folha.cell(row=numero, column=coluna)
            celula.fill = amarelo
            celula.border = borda
            celula.font = negrito if numero == total_linha else normal
            celula.alignment = esquerda if coluna == 5 else centro

    larguras = {1: 42, 2: 30, 3: 14, 4: 12, 5: 32}
    for coluna in range(1, colunas_totais + 1):
        folha.column_dimensions[get_column_letter(coluna)].width = larguras.get(coluna, 9)

    saida = BytesIO()
    planilha.save(saida)
    return saida.getvalue()
