"""Estilos compartilhados pelos relatorios .xlsx das atividades."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FINO = Side(style="thin", color="000000")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
AZUL = PatternFill("solid", fgColor="DDEBF7")
AMARELO = PatternFill("solid", fgColor="FFFF00")
NEGRITO = Font(name="Calibri", size=11, bold=True)
NORMAL = Font(name="Calibri", size=11)
CENTRO = Alignment(horizontal="center", vertical="center")
ESQUERDA = Alignment(horizontal="left", vertical="center")
