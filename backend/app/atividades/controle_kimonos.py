"""Relatorio "CONTROLE DE ENTREGA DE KIMONOS" da atividade de entrega de kimono.

Monta as linhas (nome, turma, tamanho), a tabela de quantidades por turma e
tamanho e a planilha .xlsx no mesmo layout do controle de faixas.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ._estilo_xlsx import AMARELO, AZUL, BORDA, CENTRO, ESQUERDA, NEGRITO, NORMAL

TITULO = "CONTROLE DE ENTREGA DE KIMONOS"
TAMANHOS = ["PP", "P", "N2", "M", "G", "N4", "GG"]
SEM_TURMA = "SEM TURMA"
LINHAS_MINIMAS = 14


def normalizar_turma(texto):
    return str(texto).strip().upper() if texto else SEM_TURMA


def normalizar_tamanho(texto):
    return str(texto).strip().upper() if texto else ""


def montar_linhas(alunos):
    return [
        {
            "nome": aluno.nome,
            "turma": normalizar_turma(getattr(aluno, "turma", None)),
            "tamanho": normalizar_tamanho(getattr(aluno, "tamanho_kimono", None)),
        }
        for aluno in alunos
    ]


def calcular_quantidades(linhas):
    turmas = sorted({linha["turma"] for linha in linhas}) or [SEM_TURMA]
    matriz = {turma: {tamanho: 0 for tamanho in TAMANHOS} for turma in turmas}
    for linha in linhas:
        if linha["turma"] in matriz and linha["tamanho"] in TAMANHOS:
            matriz[linha["turma"]][linha["tamanho"]] += 1
    por_turma = [
        {
            "turma": turma,
            "tamanhos": [matriz[turma][t] for t in TAMANHOS],
            "total": sum(matriz[turma].values()),
        }
        for turma in turmas
    ]
    totais = [sum(matriz[t][tam] for t in turmas) for tam in TAMANHOS]
    return {
        "tamanhos": TAMANHOS,
        "turmas": por_turma,
        "totais": totais,
        "total_geral": sum(totais),
    }


def gerar_relatorio(alunos):
    linhas = montar_linhas(alunos)
    return {"titulo": TITULO, "linhas": linhas, "quantidades": calcular_quantidades(linhas)}


def gerar_xlsx(alunos):
    linhas = montar_linhas(alunos)
    quantidades = calcular_quantidades(linhas)
    turmas = quantidades["turmas"]
    ultima_lista = 2 + max(len(linhas), LINHAS_MINIMAS)
    total_linha = 3 + len(turmas)

    planilha = Workbook()
    folha = planilha.active
    folha.title = "Controle de kimonos"

    colunas_totais = 4 + len(TAMANHOS) + 1
    ultima_coluna = get_column_letter(colunas_totais)
    folha.merge_cells(f"A1:{ultima_coluna}1")
    folha["A1"] = TITULO
    folha["A1"].font = NEGRITO
    folha["A1"].alignment = CENTRO
    for coluna in range(1, colunas_totais + 1):
        folha.cell(row=1, column=coluna).border = BORDA

    cabecalhos = ["NOME", "TURMA", "TAMANHO", "QUANTIDADES", *TAMANHOS, "Total"]
    for indice, texto in enumerate(cabecalhos, start=1):
        celula = folha.cell(row=2, column=indice, value=texto)
        celula.font = NEGRITO
        celula.alignment = CENTRO
        celula.border = BORDA

    for numero in range(3, ultima_lista + 1):
        for coluna in range(1, 4):
            celula = folha.cell(row=numero, column=coluna)
            celula.border = BORDA
            celula.font = NORMAL
            celula.alignment = ESQUERDA if coluna == 1 else CENTRO
            if coluna in (2, 3):
                celula.fill = AZUL
    for indice, linha in enumerate(linhas):
        numero = 3 + indice
        folha.cell(row=numero, column=1, value=linha["nome"])
        folha.cell(row=numero, column=2, value=linha["turma"] or None)
        folha.cell(row=numero, column=3, value=linha["tamanho"] or None)

    intervalo_turma = f"$B$3:$B${ultima_lista}"
    intervalo_tamanho = f"$C$3:$C${ultima_lista}"
    for indice, resumo in enumerate(turmas):
        numero = 3 + indice
        folha.cell(row=numero, column=4, value=resumo["turma"])
        for posicao, tamanho in enumerate(TAMANHOS):
            coluna = 5 + posicao
            folha.cell(
                row=numero,
                column=coluna,
                value=f'=COUNTIFS({intervalo_turma},$D{numero},{intervalo_tamanho},{get_column_letter(coluna)}$2)',
            )
        primeira, ultima = get_column_letter(5), get_column_letter(4 + len(TAMANHOS))
        folha.cell(row=numero, column=colunas_totais, value=f"=SUM({primeira}{numero}:{ultima}{numero})")

    folha.cell(row=total_linha, column=4, value="TOTAL:")
    for coluna in range(5, colunas_totais + 1):
        letra = get_column_letter(coluna)
        folha.cell(row=total_linha, column=coluna, value=f"=SUM({letra}3:{letra}{total_linha - 1})")

    for numero in range(3, total_linha + 1):
        for coluna in range(4, colunas_totais + 1):
            celula = folha.cell(row=numero, column=coluna)
            celula.fill = AMARELO
            celula.border = BORDA
            celula.font = NEGRITO if numero == total_linha else NORMAL
            celula.alignment = ESQUERDA if coluna == 4 else CENTRO

    larguras = {1: 42, 2: 28, 3: 14, 4: 32}
    for coluna in range(1, colunas_totais + 1):
        folha.column_dimensions[get_column_letter(coluna)].width = larguras.get(coluna, 9)

    saida = BytesIO()
    planilha.save(saida)
    return saida.getvalue()
